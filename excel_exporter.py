import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import database

def generate_excel_report():
    """
    Retrieves the complete flattened relational dataset from the database
    and builds a highly-styled executive Excel workbook in memory.
    """
    df = database.get_unified_search_df()
    
    # Fill empty values for professional presentation
    df['Responsável'] = df['Responsável'].fillna('Almoxarifado (Disponível)')
    df['Localização'] = df['Localização'].fillna('Almoxarifado (Geral)')
    df['Mesa'] = df['Mesa'].fillna('-')
    
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Ativos GAP', index=False)
        
        # Access openpyxl objects to inject corporate design tokens
        workbook = writer.book
        worksheet = writer.sheets['Ativos GAP']
        
        # Show gridlines (very important for enterprise spreadsheets)
        worksheet.views.sheetView[0].showGridLines = True
        
        # Design Tokens
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Navy Blue
        
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center')
        
        # Header formatting
        for col_num, col_name in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            
            # Dynamic Width Calculation (vectorized & safe against NaNs)
            max_len = max(
                df[col_name].fillna("").astype(str).str.len().max(),
                len(col_name)
            ) + 4
            worksheet.column_dimensions[worksheet.cell(row=1, column=col_num).column_letter].width = min(max_len, 45)
            
        # Cell and row formatting (alternating rows, thin borders, alignments)
        data_font = Font(name='Segoe UI', size=10)
        border_side = Side(border_style="thin", color="E5E7EB")
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        zebra_fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid') # soft gray
        
        for row_idx in range(2, len(df) + 2):
            fill = zebra_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = cell_border
                if fill.fill_type:
                    cell.fill = fill
                
                # Alignments based on content
                col_name = df.columns[col_idx - 1]
                if col_name in ['Código', 'Status', 'Mesa']:
                    cell.alignment = center_alignment
                else:
                    cell.alignment = left_alignment
                    
        # Set Header row height
        worksheet.row_dimensions[1].height = 28
        
    return output.getvalue()
